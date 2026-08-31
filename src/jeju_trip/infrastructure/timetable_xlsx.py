"""공식 제주 버스 XLSX raw bundle을 안전하게 읽는 importer 전용 경계."""

from __future__ import annotations

import hashlib
import io
import json
import re
import warnings
import zipfile
from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import PurePosixPath
from typing import BinaryIO


class TimetableWorkbookRejected(ValueError):
    """원본 계보나 workbook 안전 조건을 충족하지 못한 경우."""


@dataclass(frozen=True)
class WorkbookCellRow:
    sheet_name: str
    row_number: int
    values: tuple[object, ...]


@dataclass(frozen=True)
class OfficialTimetableRawBundle:
    """공식 raw ZIP에서 검증된 workbook 값과 lineage manifest."""

    workbook_rows: dict[str, tuple[WorkbookCellRow, ...]]
    mapping_manifest: dict[str, object]
    service_day_manifest: dict[str, object]
    source_references: tuple[str, ...]


OFFICIAL_WORKBOOK_CHECKSUMS = {
    "201": frozenset(
        {
            "d3a6b9db15f0786a1ac1ab7865a7762f51b7521ef69e9844b1ca6e19b0e2b04e",
            "75806195a891543ea6ec2dd5255dabd14613753e862554af313a9efcf13f0efe",
        }
    ),
    "211-212": frozenset(
        {
            "2e0e13c4d980b7edaaa7540e2ffc013bdb3d43e9e19c7a46caccc66584c615eb",
            "e6a8479b00536dad2b2e81e16bd0f02547fe7c463f9be81ba84641eb18aafbaa",
        }
    ),
}
OFFICIAL_SNAPSHOT_CHECKSUM_MANIFEST_SHA256 = {
    "jeju-bis-2026-08-17-234": (
        "b298bdb99b4597157a9a5d1bb2d3c116478ed10055988c6dea4fc3795ffb8481"
    )
}
_SOURCE_CLOCK = re.compile(r"^\s*(\d{1,2}):(\d{2})")


def sha256_bytes(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def _canonical_json_sha256(value: object) -> str:
    payload = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return sha256_bytes(payload)


def _validate_snapshot_checksum_manifest(
    snapshot_id: object,
    checksum_manifest: dict[str, object],
) -> bool:
    """알려진 전역 snapshot의 전체 workbook checksum 집합만 승인한다."""

    if snapshot_id is None:
        return False
    if not isinstance(snapshot_id, str) or snapshot_id not in (
        OFFICIAL_SNAPSHOT_CHECKSUM_MANIFEST_SHA256
    ):
        raise TimetableWorkbookRejected("TIMETABLE_OFFICIAL_SNAPSHOT_NOT_PINNED")
    if checksum_manifest.get("algorithm") != "SHA256":
        raise TimetableWorkbookRejected("TIMETABLE_CHECKSUM_ALGORITHM_INVALID")
    if _canonical_json_sha256(checksum_manifest) != (
        OFFICIAL_SNAPSHOT_CHECKSUM_MANIFEST_SHA256[snapshot_id]
    ):
        raise TimetableWorkbookRejected("TIMETABLE_OFFICIAL_SNAPSHOT_CHECKSUM_MISMATCH")
    return True


def _source_clocks(values: tuple[object, ...]) -> tuple[tuple[int, str], ...]:
    result: list[tuple[int, str]] = []
    for index, value in enumerate(values):
        text = str(value or "")
        if len(re.findall(r"(?<!\d)\d{1,2}:\d{2}(?!\d)", text)) != 1:
            continue
        matched = _SOURCE_CLOCK.match(text)
        if matched is None:
            continue
        hour = int(matched.group(1))
        minute = int(matched.group(2))
        if hour <= 23 and minute <= 59:
            result.append((index, f"{hour:02d}{minute:02d}"))
    return tuple(result)


def _service_clock_on_source_row(value: object) -> str | None:
    text = str(value or "")
    if not re.fullmatch(r"\d{4}", text):
        return None
    hour = int(text[:2])
    minute = int(text[2:])
    if hour > 71 or minute > 59:
        return None
    return f"{hour % 24:02d}{minute:02d}"


def _validate_snapshot_mapping_lineage(
    workbook_rows: Mapping[str, tuple[WorkbookCellRow, ...]],
    mapping_manifest: dict[str, object],
) -> None:
    """전역 derived manifest의 시각·정류장 순서가 pinned workbook 행과 연결되는지 확인한다."""

    raw_route_stops = mapping_manifest.get("route_stops")
    trip_rows = mapping_manifest.get("trip_rows")
    if not isinstance(raw_route_stops, list) or not isinstance(trip_rows, list):
        raise TimetableWorkbookRejected("TIMETABLE_RAW_LINEAGE_INVALID")
    pattern_stops: dict[str, set[tuple[str, int]]] = {}
    for row in raw_route_stops:
        if not isinstance(row, dict):
            raise TimetableWorkbookRejected("TIMETABLE_ROUTE_STOP_LINEAGE_INVALID")
        pattern = row.get("provider_route_pattern_id")
        stop = row.get("provider_stop_id")
        sequence = row.get("route_sequence")
        if (
            not isinstance(pattern, str)
            or not isinstance(stop, str)
            or not isinstance(sequence, int)
            or sequence <= 0
            or (stop, sequence) in pattern_stops.setdefault(pattern, set())
        ):
            raise TimetableWorkbookRejected("TIMETABLE_ROUTE_STOP_LINEAGE_INVALID")
        pattern_stops[pattern].add((stop, sequence))

    source_rows = {
        (filename, row.sheet_name, row.row_number): row.values
        for filename, rows in workbook_rows.items()
        for row in rows
    }
    for trip in trip_rows:
        if not isinstance(trip, dict):
            raise TimetableWorkbookRejected("TIMETABLE_MAPPING_ROW_INVALID")
        workbook = trip.get("workbook")
        sheet_name = trip.get("sheet_name")
        row_number = trip.get("source_row_number")
        patterns = trip.get("provider_route_pattern_ids")
        stop_times = trip.get("major_stop_times")
        if (
            not isinstance(workbook, str)
            or not isinstance(sheet_name, str)
            or not isinstance(row_number, int)
            or not isinstance(patterns, list)
            or len(patterns) != 1
            or not isinstance(patterns[0], str)
            or not isinstance(stop_times, list)
            or len(stop_times) < 2
        ):
            raise TimetableWorkbookRejected("TIMETABLE_MAPPING_ROW_LINEAGE_INVALID")
        values = source_rows.get((workbook, sheet_name, row_number))
        if values is None:
            raise TimetableWorkbookRejected("TIMETABLE_SOURCE_ROW_MISSING")
        route = pattern_stops.get(patterns[0])
        if route is None:
            raise TimetableWorkbookRejected("TIMETABLE_ROUTE_PATTERN_LINEAGE_MISSING")

        source_clocks = _source_clocks(values)
        last_source_index = -1
        last_route_sequence = 0
        for stop_time in stop_times:
            if not isinstance(stop_time, dict):
                raise TimetableWorkbookRejected("TIMETABLE_MAJOR_STOP_TIME_INVALID")
            stop = stop_time.get("provider_stop_id")
            route_sequence = stop_time.get("route_sequence")
            arrival = stop_time.get("arrival_time")
            departure = stop_time.get("departure_time")
            source_clock = _service_clock_on_source_row(arrival)
            if (
                not isinstance(stop, str)
                or not isinstance(route_sequence, int)
                or (stop, route_sequence) not in route
                or route_sequence <= last_route_sequence
                or arrival != departure
                or source_clock is None
            ):
                raise TimetableWorkbookRejected("TIMETABLE_MAJOR_STOP_TIME_LINEAGE_INVALID")
            matching_source = next(
                (
                    index
                    for index, clock in source_clocks
                    if index > last_source_index and clock == source_clock
                ),
                None,
            )
            if matching_source is None:
                raise TimetableWorkbookRejected("TIMETABLE_SOURCE_CLOCK_MISMATCH")
            last_source_index = matching_source
            last_route_sequence = route_sequence


def _validate_xlsx_container(payload: bytes) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(payload))
    except zipfile.BadZipFile as error:
        raise TimetableWorkbookRejected("TIMETABLE_XLSX_INVALID") from error
    total_uncompressed = 0
    for member in archive.infolist():
        path = PurePosixPath(member.filename)
        if path.is_absolute() or ".." in path.parts:
            raise TimetableWorkbookRejected("TIMETABLE_XLSX_PATH_TRAVERSAL")
        total_uncompressed += member.file_size
        ratio = member.file_size / max(1, member.compress_size)
        if ratio > 100:
            raise TimetableWorkbookRejected("TIMETABLE_XLSX_COMPRESSION_RATIO_EXCEEDED")
        lowered = member.filename.lower()
        if "externallinks/" in lowered:
            raise TimetableWorkbookRejected("TIMETABLE_XLSX_EXTERNAL_LINK_REJECTED")
        if lowered.endswith("vbaproject.bin"):
            raise TimetableWorkbookRejected("TIMETABLE_XLSX_MACRO_REJECTED")
        if lowered.startswith("xl/worksheets/") and lowered.endswith(".xml"):
            xml = archive.read(member)
            if b"<f" in xml:
                raise TimetableWorkbookRejected("TIMETABLE_XLSX_FORMULA_REJECTED")
    if total_uncompressed > 100 * 1024 * 1024:
        raise TimetableWorkbookRejected("TIMETABLE_XLSX_UNCOMPRESSED_SIZE_EXCEEDED")


def read_official_workbook(
    stream: BinaryIO,
    *,
    expected_sha256: str,
    required_sheet_tokens: tuple[str, ...],
) -> tuple[WorkbookCellRow, ...]:
    """checksum과 sheet 노선·방향 token을 확인하고 값만 메모리에 반환한다."""

    payload = stream.read()
    if sha256_bytes(payload) != expected_sha256:
        raise TimetableWorkbookRejected("TIMETABLE_XLSX_CHECKSUM_MISMATCH")
    _validate_xlsx_container(payload)

    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - importer dependency 누락 환경
        raise TimetableWorkbookRejected("OPENPYXL_IMPORTER_DEPENDENCY_MISSING") from error

    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style, apply openpyxl's default",
            category=UserWarning,
        )
        workbook = load_workbook(
            io.BytesIO(payload),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    try:
        if not all(
            any(token in sheet_name for sheet_name in workbook.sheetnames)
            for token in required_sheet_tokens
        ):
            raise TimetableWorkbookRejected("TIMETABLE_XLSX_SHEET_IDENTITY_MISMATCH")
        rows: list[WorkbookCellRow] = []
        for sheet in workbook.worksheets:
            # 제주 BIS workbook은 실제 셀이 있어도 dimension을 A1로 기록하는 경우가 있다.
            # read_only mode는 이 값을 신뢰하므로 원본을 변경하지 않고 범위를 재탐색한다.
            sheet.reset_dimensions()
            for number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
                if any(value not in (None, "") for value in cells):
                    rows.append(WorkbookCellRow(sheet.title, number, tuple(cells)))
        return tuple(rows)
    finally:
        workbook.close()


def require_holiday_service_evidence(source_references: tuple[str, ...]) -> None:
    """공휴일 적용을 공식 HTTPS 근거 없이 생성하지 않는다."""

    if not source_references or any(not ref.startswith("https://") for ref in source_references):
        raise TimetableWorkbookRejected("BUS_SERVICE_DAY_UNVERIFIED")


def service_day_source_references(manifest: dict[str, object]) -> tuple[str, ...]:
    """평일 근거와 선택적 공휴일 근거를 분리해 검증한다."""

    weekday = manifest.get("weekday_source_references")
    if (
        not isinstance(weekday, list)
        or not weekday
        or not all(isinstance(ref, str) and ref.startswith("https://") for ref in weekday)
    ):
        raise TimetableWorkbookRejected("BUS_WEEKDAY_SERVICE_DAY_UNVERIFIED")
    holiday = manifest.get("holiday_source_references", [])
    if not isinstance(holiday, list) or not all(isinstance(ref, str) for ref in holiday):
        raise TimetableWorkbookRejected("BUS_SERVICE_DAY_UNVERIFIED")
    if holiday:
        require_holiday_service_evidence(tuple(holiday))
    return tuple(dict.fromkeys((*weekday, *holiday)))


def load_official_timetable_zip(stream: BinaryIO) -> OfficialTimetableRawBundle:
    """공식 XLSX·checksum·mapping·service-day manifest를 하나의 raw ZIP에서 검증한다."""

    payload = stream.read()
    try:
        archive = zipfile.ZipFile(io.BytesIO(payload))
    except zipfile.BadZipFile as error:
        raise TimetableWorkbookRejected("TIMETABLE_RAW_ZIP_INVALID") from error
    names = set(archive.namelist())
    required = {
        "checksum-manifest.json",
        "mapping-manifest.json",
        "service-day-manifest.json",
    }
    if not required <= names:
        raise TimetableWorkbookRejected("TIMETABLE_RAW_MANIFEST_MISSING")
    if any(
        PurePosixPath(name).is_absolute() or ".." in PurePosixPath(name).parts for name in names
    ):
        raise TimetableWorkbookRejected("TIMETABLE_RAW_PATH_TRAVERSAL")
    try:
        checksum_manifest = json.loads(archive.read("checksum-manifest.json"))
        mapping_manifest = json.loads(archive.read("mapping-manifest.json"))
        service_day_manifest = json.loads(archive.read("service-day-manifest.json"))
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise TimetableWorkbookRejected("TIMETABLE_RAW_MANIFEST_INVALID") from error
    manifests = (checksum_manifest, mapping_manifest, service_day_manifest)
    if not all(isinstance(item, dict) for item in manifests):
        raise TimetableWorkbookRejected("TIMETABLE_RAW_MANIFEST_ROOT_INVALID")

    checksums = checksum_manifest.get("files")
    workbook_specs = mapping_manifest.get("workbooks")
    trip_rows = mapping_manifest.get("trip_rows")
    if not isinstance(checksums, dict) or not isinstance(workbook_specs, dict):
        raise TimetableWorkbookRejected("TIMETABLE_RAW_LINEAGE_INVALID")
    if not isinstance(trip_rows, list) or not trip_rows:
        raise TimetableWorkbookRejected("TIMETABLE_MAPPING_ROWS_EMPTY")
    snapshot_pinned = _validate_snapshot_checksum_manifest(
        mapping_manifest.get("official_snapshot_id"), checksum_manifest
    )
    if snapshot_pinned:
        if set(workbook_specs) != set(checksums):
            raise TimetableWorkbookRejected("TIMETABLE_RAW_LINEAGE_INVALID")
        if names != required | set(workbook_specs):
            raise TimetableWorkbookRejected("TIMETABLE_RAW_UNTRACKED_MEMBER")

    workbook_rows: dict[str, tuple[WorkbookCellRow, ...]] = {}
    matched_shapes: set[str] = set()
    for filename, spec in workbook_specs.items():
        if not isinstance(filename, str) or not filename.lower().endswith(".xlsx"):
            raise TimetableWorkbookRejected("TIMETABLE_WORKBOOK_NAME_INVALID")
        if filename not in names or not isinstance(spec, dict):
            raise TimetableWorkbookRejected("TIMETABLE_WORKBOOK_MISSING")
        expected = checksums.get(filename)
        tokens = spec.get("required_sheet_tokens")
        route_group = spec.get("route_group")
        if (
            not isinstance(expected, str)
            or len(expected) != 64
            or not isinstance(tokens, list)
            or not all(isinstance(token, str) and token for token in tokens)
            or (not snapshot_pinned and route_group not in OFFICIAL_WORKBOOK_CHECKSUMS)
        ):
            raise TimetableWorkbookRejected("TIMETABLE_WORKBOOK_MANIFEST_INVALID")
        if not snapshot_pinned and expected not in OFFICIAL_WORKBOOK_CHECKSUMS[str(route_group)]:
            raise TimetableWorkbookRejected("TIMETABLE_OFFICIAL_CHECKSUM_NOT_PINNED")
        workbook_rows[filename] = read_official_workbook(
            io.BytesIO(archive.read(filename)),
            expected_sha256=expected,
            required_sheet_tokens=tuple(tokens),
        )
        matched_shapes.add(str(route_group))
    if not snapshot_pinned and matched_shapes != set(OFFICIAL_WORKBOOK_CHECKSUMS):
        raise TimetableWorkbookRejected("TIMETABLE_REQUIRED_WORKBOOKS_MISSING")

    if snapshot_pinned:
        _validate_snapshot_mapping_lineage(workbook_rows, mapping_manifest)

    for row in trip_rows:
        if not isinstance(row, dict):
            raise TimetableWorkbookRejected("TIMETABLE_MAPPING_ROW_INVALID")
        patterns = row.get("provider_route_pattern_ids")
        stop_times = row.get("major_stop_times")
        if not isinstance(patterns, list) or len(patterns) != 1:
            raise TimetableWorkbookRejected("TIMETABLE_ROUTE_PATTERN_NOT_UNIQUE")
        if not isinstance(stop_times, list) or len(stop_times) < 2:
            raise TimetableWorkbookRejected("TIMETABLE_MAJOR_STOP_TIMES_INSUFFICIENT")

    references = service_day_source_references(service_day_manifest)
    return OfficialTimetableRawBundle(
        workbook_rows=workbook_rows,
        mapping_manifest=mapping_manifest,
        service_day_manifest=service_day_manifest,
        source_references=references,
    )
